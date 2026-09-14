"""
resume_parser.py

Turns a raw resume (PDF, DOCX, or pasted text) into structured data:
contact info, detected sections, skills found, and some rough counts
that the analyzer uses for scoring.

Extraction is intentionally simple and rule/regex based so the whole
project runs without any paid AI API.
"""

import io
import re

import pdfplumber
from docx import Document

from job_roles import ALL_KNOWN_SKILLS

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")
PHONE_RE = re.compile(r"(\+?\d{1,3}[-.\s]?)?\(?\d{3,5}\)?[-.\s]?\d{3}[-.\s]?\d{3,4}")
LINKEDIN_RE = re.compile(r"(linkedin\.com/[^\s,)]+)", re.IGNORECASE)
GITHUB_RE = re.compile(r"(github\.com/[^\s,)]+)", re.IGNORECASE)

SECTION_HEADINGS = {
    "education": ["education", "academic background", "qualification"],
    "experience": ["experience", "work experience", "employment", "internship"],
    "projects": ["projects", "academic projects", "personal projects"],
    "skills": ["skills", "technical skills", "core competencies"],
    "summary": ["summary", "objective", "profile"],
    "certifications": ["certifications", "certificates", "licenses"],
}

ACTION_VERBS = [
    "developed", "built", "designed", "implemented", "led", "managed",
    "created", "improved", "optimized", "analyzed", "automated",
    "launched", "reduced", "increased", "delivered", "collaborated",
]


class ResumeParseError(Exception):
    """Raised when a resume file can't be read or is unusable."""


def extract_text_from_pdf(file_bytes):
    try:
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                text_parts.append(page_text)
        text = "\n".join(text_parts).strip()
        if not text:
            raise ResumeParseError(
                "We couldn't find any readable text in this PDF. "
                "It may be a scanned image — try pasting the text instead."
            )
        return text
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError(f"This PDF could not be read ({exc}).") from exc


def extract_text_from_docx(file_bytes):
    try:
        document = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)
        text = "\n".join(p for p in paragraphs if p is not None).strip()
        if not text:
            raise ResumeParseError("This DOCX file appears to be empty.")
        return text
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError(f"This DOCX file could not be read ({exc}).") from exc


def extract_text(filename, file_bytes):
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    if lower.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    raise ResumeParseError("Unsupported file type. Please upload a PDF or DOCX file.")


def find_contact_info(text):
    email_match = EMAIL_RE.search(text)
    phone_match = PHONE_RE.search(text)
    linkedin_match = LINKEDIN_RE.search(text)
    github_match = GITHUB_RE.search(text)

    # Name heuristic: first non-empty line that isn't an email/phone/url
    # and reads like a short human name (2-4 words, no digits).
    name = None
    for line in text.splitlines()[:8]:
        candidate = line.strip()
        if not candidate or len(candidate) > 40:
            continue
        if EMAIL_RE.search(candidate) or PHONE_RE.search(candidate):
            continue
        words = candidate.split()
        if 1 < len(words) <= 4 and all(w.replace(".", "").isalpha() for w in words):
            name = candidate
            break

    # Location heuristic: prefer the contact-details line (the one with the
    # email/phone), split it on common separators, and pick the short
    # remaining segment that looks like "City, ST" rather than a sentence.
    location = None
    contact_line = None
    for line in text.splitlines()[:10]:
        if EMAIL_RE.search(line) or PHONE_RE.search(line):
            contact_line = line
            break

    def _looks_like_place(segment):
        segment = segment.strip()
        if not segment or len(segment) > 40 or segment.endswith("."):
            return False
        words = segment.split()
        if len(words) > 4:
            return False
        return True

    if contact_line:
        for part in re.split(r"[|•·]", contact_line):
            part = part.strip()
            if part and not EMAIL_RE.search(part) and not PHONE_RE.search(part) \
                    and not LINKEDIN_RE.search(part) and not GITHUB_RE.search(part) \
                    and _looks_like_place(part):
                location = part
                break

    if not location:
        for line in text.splitlines()[:10]:
            candidate = line.strip()
            if "," in candidate and not EMAIL_RE.search(candidate) and not any(c.isdigit() for c in candidate) \
                    and _looks_like_place(candidate):
                location = candidate
                break

    return {
        "name": name or "Not detected",
        "email": email_match.group(0) if email_match else None,
        "phone": phone_match.group(0) if phone_match else None,
        "location": location,
        "linkedin": linkedin_match.group(1) if linkedin_match else None,
        "github": github_match.group(1) if github_match else None,
    }


def detect_sections(text):
    """Return {section_name: bool_present} based on heading keywords."""
    lower = text.lower()
    found = {}
    for section, keywords in SECTION_HEADINGS.items():
        found[section] = any(kw in lower for kw in keywords)
    return found


def extract_skills(text):
    lower = text.lower()
    found = []
    for skill in ALL_KNOWN_SKILLS:
        pattern = r"(?<![a-z0-9])" + re.escape(skill) + r"(?![a-z0-9])"
        if re.search(pattern, lower):
            found.append(skill)
    return sorted(set(found))


def count_action_verbs(text):
    lower = text.lower()
    return sum(len(re.findall(r"\b" + verb + r"\b", lower)) for verb in ACTION_VERBS)


def count_quantifiable_achievements(text):
    # Looks for bullet-style lines containing a number/percentage,
    # a rough proxy for "measurable achievements".
    lines = text.splitlines()
    hits = 0
    for line in lines:
        if re.search(r"\d", line) and re.search(r"%|\$|\bhours?\b|\busers?\b|\bcustomers?\b|\d+x\b", line, re.IGNORECASE):
            hits += 1
    return hits


def parse_resume(text):
    """Main entry point: turn raw resume text into a structured dict."""
    cleaned = text.strip()
    if len(cleaned) < 50:
        raise ResumeParseError(
            "This resume looks too short to analyze. Please provide a more complete resume."
        )

    contact = find_contact_info(cleaned)
    sections = detect_sections(cleaned)
    skills = extract_skills(cleaned)
    action_verbs = count_action_verbs(cleaned)
    quantifiable = count_quantifiable_achievements(cleaned)
    word_count = len(cleaned.split())

    return {
        "raw_text": cleaned,
        "contact": contact,
        "sections": sections,
        "skills": skills,
        "action_verb_count": action_verbs,
        "quantifiable_achievements": quantifiable,
        "word_count": word_count,
    }


DEMO_RESUME_TEXT = """Alex Johnson
alex.johnson@email.com | +1 555-201-4488 | Austin, TX
linkedin.com/in/alexjohnson | github.com/alexjohnson

Summary
Data Analyst intern with hands-on experience turning raw business data into
clear, actionable insight using Python and SQL.

Education
B.Tech Computer Science, University of Texas, 2022-2026

Skills
Python, Excel, SQL, Power BI, Pandas, Statistics, Data Visualization

Experience
Data Analyst Intern, BrightPath Retail (Jun 2025 - Aug 2025)
Built a sales dashboard in Power BI that reduced weekly reporting time by 30%.
Automated a Python data-cleaning pipeline used by 5 analysts.

Projects
Sales Dashboard - Analyzed 2 years of transaction data to identify seasonal trends.
Customer Churn Analysis - Built a churn model in Python achieving 85% accuracy.

Certifications
Google Data Analytics Certificate
"""
